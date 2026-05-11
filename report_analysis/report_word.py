"""
Generate a Word (.docx) compliance report using python-docx.
"""

import os

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


_OUTCOME_BADGE = {
    "PASS":    ("PASS",    RGBColor(0x22, 0xC5, 0x5E)),
    "WARNING": ("WARNING", RGBColor(0xF5, 0x9E, 0x0B)),
    "FAIL":    ("FAIL",    RGBColor(0xEF, 0x44, 0x44)),
    "N/A":     ("N/A",     RGBColor(0x6B, 0x72, 0x80)),
}

_PRIMARY   = RGBColor(0x00, 0x5B, 0x8E)
_TXT       = RGBColor(0x1A, 0x1A, 0x2E)
_TXT2      = RGBColor(0x6B, 0x72, 0x80)
_ERROR     = RGBColor(0xEF, 0x44, 0x44)
_COMPLIANT = RGBColor(0x22, 0xC5, 0x5E)


def _set_run_color(run, color: RGBColor):
    run.font.color.rgb = color


def _heading(doc, text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    if color:
        for run in p.runs:
            _set_run_color(run, color)
    return p


def _para(doc, text, bold=False, italic=False, color=None, size=None, alignment=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        _set_run_color(run, color)
    if size:
        run.font.size = Pt(size)
    if alignment:
        p.alignment = alignment
    return p


def _kv_row(doc, key, value, key_color=None, val_color=None):
    p = doc.add_paragraph()
    k = p.add_run(f"{key}: ")
    k.bold = True
    if key_color:
        _set_run_color(k, key_color)
    v = p.add_run(value)
    if val_color:
        _set_run_color(v, val_color)
    return p


def _para_spacing(p, before, after):
    p.paragraph_format.space_before = Pt(before / 20)
    p.paragraph_format.space_after = Pt(after / 20)


def _divider(doc):
    p = doc.add_paragraph()
    _para_spacing(p, 60, 60)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "F5A623")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def generate_word_report(report_data: dict, output_dir: str = "reports/") -> str:
    """
    Generate a Word document from report_data and save it to output_dir.
    Returns the absolute path of the saved file.
    """
    os.makedirs(output_dir, exist_ok=True)

    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.2)
        section.right_margin  = Inches(1.2)

    # ── Cover block ──────────────────────────────────────────────────────────
    org = doc.add_paragraph()
    org_run = org.add_run("Sunshine Coast Council")
    org_run.bold = True
    org_run.font.size = Pt(11)
    _set_run_color(org_run, _PRIMARY)

    title_p = doc.add_heading(f"Compliance Report #{report_data['report_id']}", level=0)
    for run in title_p.runs:
        _set_run_color(run, _PRIMARY)

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Software Contract Compliance Screening")
    sub_run.font.size = Pt(12)
    _set_run_color(sub_run, _TXT2)

    _divider(doc)

    # ── Metadata table ───────────────────────────────────────────────────────
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.style = "Table Grid"
    meta_fields = [
        ("Generated",     report_data["generated_at"]),
        ("Contract",      report_data["contract_name"]),
        ("Vendor / Party",report_data["vendor_name"]),
        ("Analyst",       report_data["analyst"]),
    ]
    for i, (label, value) in enumerate(meta_fields):
        row = meta_table.rows[i]
        key_cell  = row.cells[0]
        val_cell  = row.cells[1]
        key_run = key_cell.paragraphs[0].add_run(label)
        key_run.bold = True
        key_run.font.size = Pt(10)
        _set_run_color(key_run, _TXT2)
        val_run = val_cell.paragraphs[0].add_run(value or "—")
        val_run.font.size = Pt(10)
        _set_run_color(val_run, _TXT)

    doc.add_paragraph()

    # ── Summary ───────────────────────────────────────────────────────────────
    _heading(doc, "Executive Summary", level=1, color=_PRIMARY)

    s = report_data["summary"]

    stats_table = doc.add_table(rows=2, cols=5)
    stats_table.style = "Table Grid"
    headers = ["Total Rules", "Passed", "Warnings", "Failed", "Critical Failures"]
    values  = [
        str(s["total_rules"]),
        str(s["passed"]),
        str(s["warnings"]),
        str(s["failed"]),
        str(s["critical_failures"]),
    ]
    value_colors = [_TXT, _COMPLIANT, RGBColor(0xF5, 0x9E, 0x0B), _ERROR, _ERROR]
    for col_idx, (h, v, vc) in enumerate(zip(headers, values, value_colors)):
        hdr_run = stats_table.rows[0].cells[col_idx].paragraphs[0].add_run(h)
        hdr_run.bold = True
        hdr_run.font.size = Pt(9)
        _set_run_color(hdr_run, _TXT2)
        val_run = stats_table.rows[1].cells[col_idx].paragraphs[0].add_run(v)
        val_run.bold = True
        val_run.font.size = Pt(13)
        _set_run_color(val_run, vc)

    doc.add_paragraph()
    _divider(doc)

    # ── Findings ──────────────────────────────────────────────────────────────
    _heading(doc, "Detailed Findings", level=1, color=_PRIMARY)

    for finding in report_data["findings"]:
        outcome      = finding.get("outcome", "N/A")
        badge_text, badge_color = _OUTCOME_BADGE.get(outcome, ("N/A", _TXT2))
        is_critical  = finding.get("critical", False)
        rule_id      = finding.get("rule_id", "—")
        rule_title   = finding.get("rule_title", "")

        # Finding heading
        hdr_p = doc.add_heading("", level=2)
        hdr_p.clear()
        id_run = hdr_p.add_run(rule_id)
        id_run.bold = True
        id_run.font.size = Pt(12)
        _set_run_color(id_run, _PRIMARY)
        if rule_title:
            sep_run = hdr_p.add_run(f"  —  {rule_title}")
            sep_run.bold = False
            sep_run.font.size = Pt(11)
            _set_run_color(sep_run, _TXT)

        # Outcome badge + critical flag
        badge_p = doc.add_paragraph()
        b_run = badge_p.add_run(f"[{badge_text}]")
        b_run.bold = True
        b_run.font.size = Pt(10)
        _set_run_color(b_run, badge_color)
        if is_critical:
            crit_run = badge_p.add_run("  [CRITICAL]")
            crit_run.bold = True
            crit_run.font.size = Pt(10)
            _set_run_color(crit_run, _ERROR)

        # Clause quoted
        clause = (finding.get("clause_quoted") or "").strip()
        if clause:
            clause_p = doc.add_paragraph()
            clause_p.paragraph_format.left_indent = Inches(0.3)
            cq_run = clause_p.add_run(f'"{clause}"')
            cq_run.italic = True
            cq_run.font.size = Pt(10)
            _set_run_color(cq_run, _TXT2)

        # Reason
        reason = (finding.get("reason") or "").strip()
        if reason:
            r_p = doc.add_paragraph()
            lbl = r_p.add_run("Finding: ")
            lbl.bold = True
            lbl.font.size = Pt(10)
            _set_run_color(lbl, _TXT)
            txt = r_p.add_run(reason)
            txt.font.size = Pt(10)
            _set_run_color(txt, _TXT)

        # Citation
        citation = (finding.get("citation") or "").strip()
        if citation:
            cite_p = doc.add_paragraph()
            cl = cite_p.add_run("Citation: ")
            cl.bold = True
            cl.font.size = Pt(9)
            _set_run_color(cl, _TXT2)
            cv = cite_p.add_run(citation)
            cv.italic = True
            cv.font.size = Pt(9)
            _set_run_color(cv, _TXT2)

        doc.add_paragraph()

    _divider(doc)

    # ── Disclaimer ────────────────────────────────────────────────────────────
    disc_p = doc.add_paragraph()
    disc_run = disc_p.add_run(
        "Disclaimer: This report is a compliance screening findings report only. "
        "It does not constitute legal advice. All findings must be reviewed by a "
        "qualified person before any procurement decision is made."
    )
    disc_run.italic = True
    disc_run.font.size = Pt(9)
    _set_run_color(disc_run, _TXT2)

    # ── Save ──────────────────────────────────────────────────────────────────
    filename = f"compliance_report_{report_data['report_id']}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return os.path.abspath(filepath)
